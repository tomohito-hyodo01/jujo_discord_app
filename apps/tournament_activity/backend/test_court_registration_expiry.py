#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
コート予約システムの利用登録の有効期限お知らせ（/api/notify/court-registration-expiry）の動作確認

DB・Discord・Google Drive のいずれにも接続せず、表の読み取り・有効期限とIDの解釈・
知らせる日の判定・文面・発行SQL・二重送信の防止・失敗時の巻き戻し・dry_run を検証する。

実行:
    cd apps/tournament_activity/backend
    python test_court_registration_expiry.py
"""

import asyncio
import os
import sys
import types
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

# Windowsの既定コンソール(cp932)でも絵文字を出せるようにする
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, 'reconfigure'):
        _stream.reconfigure(encoding='utf-8', errors='replace')

# 依存が入っていない環境（ローカルWindows等）でも動かせるよう、未導入のものだけ差し替える
for _name in ('aiomysql', 'httpx'):
    try:
        __import__(_name)
    except ImportError:  # pragma: no cover
        _stub = types.ModuleType(_name)
        _stub.DictCursor = object
        _stub.Pool = object
        _stub.AsyncClient = object
        sys.modules[_name] = _stub
try:
    from google.oauth2 import service_account  # noqa: F401
    from googleapiclient.discovery import build  # noqa: F401
    from googleapiclient.http import MediaIoBaseDownload  # noqa: F401
except ImportError:  # pragma: no cover
    for _name in ('google', 'google.oauth2', 'google.oauth2.service_account',
                  'googleapiclient', 'googleapiclient.discovery', 'googleapiclient.http'):
        sys.modules.setdefault(_name, types.ModuleType(_name))
    sys.modules['google.oauth2'].service_account = sys.modules['google.oauth2.service_account']
    sys.modules['googleapiclient.discovery'].build = None
    sys.modules['googleapiclient.http'].MediaIoBaseDownload = None

import openpyxl  # noqa: E402

import api.routers.court_registration_expiry as C  # noqa: E402

_failures = []


def check(label, actual, expected):
    if actual == expected:
        print(f"  ✅ {label}")
    else:
        _failures.append(f"{label}\n     actual  : {actual!r}\n     expected: {expected!r}")
        print(f"  ❌ {label}")


def _run(coro):
    return asyncio.run(coro)


def _workbook(rows):
    """rows: [(G列, H列, I列, M列), ...] から、1行目が見出しの表を作る（A〜F列、J〜L列は空）"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['A', 'B', 'C', 'D', 'E', 'F', '予約者名', 'カナ', '有効期限', 'テニス', '次月', '当月', 'Discord ID'])
    for g, h, i, m in rows:
        ws.append([None] * 6 + [g, h, i, None, None, None, m])
    return wb


# ===== 有効期限の読み取り =====
print("\n▼ 有効期限の読み取り")
check("表の書き方「YYYY年MM月DD日」", C.parse_expiry('2026年11月15日'), date(2026, 11, 15))
check("月日が1桁", C.parse_expiry('2026年1月5日'), date(2026, 1, 5))
check("全角数字", C.parse_expiry('２０２６年１１月１５日'), date(2026, 11, 15))
check("前後と途中の空白は無視", C.parse_expiry(' 2026年 11月 15日 '), date(2026, 11, 15))
check("日付のセル（datetime）", C.parse_expiry(datetime(2026, 11, 15, 0, 0)), date(2026, 11, 15))
check("日付のセル（date）", C.parse_expiry(date(2026, 11, 15)), date(2026, 11, 15))
check("スラッシュ区切りは読まない", C.parse_expiry('2026/11/15'), None)
check("日が無いものは読まない", C.parse_expiry('2026年11月'), None)
check("存在しない日付は読まない", C.parse_expiry('2026年2月30日'), None)
check("空は None", (C.parse_expiry(None), C.parse_expiry('')), (None, None))

# ===== Discord ID の読み取り =====
print("\n▼ Discord ID の読み取り")
check("文字のID", C.parse_discord_id('1427112485047242945'), ('1427112485047242945', 'ok'))
check("前後の空白は無視", C.parse_discord_id(' 1427112485047242945 '), ('1427112485047242945', 'ok'))
check("全角数字", C.parse_discord_id('１４２７１１２４８５０４７２４２９４５'), ('1427112485047242945', 'ok'))
check("17桁も読む", C.parse_discord_id('12345678901234567'), ('12345678901234567', 'ok'))
check("空は未記入", (C.parse_discord_id(None), C.parse_discord_id('  ')), ((None, 'missing'), (None, 'missing')))
check("数値のセル（丸まっている）は読めない", C.parse_discord_id(1.4271124850472428e18), (None, 'invalid'))
check("整数のセルも読めない", C.parse_discord_id(1427112485047242945), (None, 'invalid'))
check("メンションの書き方は読めない", C.parse_discord_id('<@1427112485047242945>'), (None, 'invalid'))
check("桁が足りないものは読めない", C.parse_discord_id('12345'), (None, 'invalid'))
check("文字が混ざると読めない", C.parse_discord_id('14271124850472429a5'), (None, 'invalid'))

# ===== 表の読み取り =====
print("\n▼ 表の読み取り")
_wb = _workbook([
    ('山田 太郎', 'ヤマダ タロウ', '2026年11月03日', '1427112485047242945'),
    (None, 'サトウ ハナコ', '2026年11月15日', None),
    ('鈴木 一郎', 'スズキ イチロウ', None, '1111111111111111111'),
    (None, None, '2026年12月01日', None),
    ('高橋 三郎', 'タカハシ サブロウ', datetime(2027, 3, 1), 1.4271124850472428e18),
    ('田中 次郎', 'タナカ ジロウ', '2026/11/15', '3333333333333333333'),
])
_accounts = C.extract_accounts(_wb)
check("有効期限が空の行と名義の無い行は読まない", [a.row_no for a in _accounts], [2, 3, 6, 7])
check("1行目（見出し）は読まない", any(a.name == '予約者名' for a in _accounts), False)
check("名義はG列（漢字）", _accounts[0].name, '山田 太郎')
check("G列が空ならH列（カタカナ）", _accounts[1].name, 'サトウ ハナコ')
check("有効期限", _accounts[0].expiry, date(2026, 11, 3))
check("ID", (_accounts[0].discord_id, _accounts[0].discord_status), ('1427112485047242945', 'ok'))
check("ID未記入", (_accounts[1].discord_id, _accounts[1].discord_status), (None, 'missing'))
check("日付のセルの有効期限", _accounts[2].expiry, date(2027, 3, 1))
check("数値のIDは読めない（元の値は残す）", (_accounts[2].discord_status, _accounts[2].discord_raw != ''), ('invalid', True))
check("読めない有効期限は None で元の値を残す", (_accounts[3].expiry, _accounts[3].expiry_raw), (None, '2026/11/15'))
check("M列が無い短い行でも読める", C._cell(('a', 'b'), C.DISCORD_ID_COL), None)
check("名義のキーは全角・半角と空白をそろえる", C.Account(2, '山田　太郎', None, '', None, 'missing', '').name_key, '山田太郎')

# ===== 知らせる日の判定 =====
print("\n▼ 知らせる日の判定")
check("1か月前（暦）", C.one_month_before(date(2026, 11, 15)), date(2026, 10, 15))
check("前月に同じ日が無ければ末日", C.one_month_before(date(2026, 3, 31)), date(2026, 2, 28))
check("うるう年の2月末", C.one_month_before(date(2028, 3, 31)), date(2028, 2, 29))
check("1月の1か月前は前年12月", C.one_month_before(date(2027, 1, 10)), date(2026, 12, 10))
_E = date(2026, 11, 15)
check("1か月前より前は知らせない", C.due_stages(_E, date(2026, 10, 14)), [])
check("1か月前の日", C.due_stages(_E, date(2026, 10, 15)), ['1m'])
check("1か月前を過ぎて未送信なら1か月前の分", C.due_stages(_E, date(2026, 10, 20)), ['1m'])
check("2週間前の日", C.due_stages(_E, date(2026, 11, 1)), ['1m', '2w'])
check("1週間前から毎日", C.due_stages(_E, date(2026, 11, 8)), ['1m', '2w', 'd2026-11-08'])
check("期限日の当日まで", C.due_stages(_E, date(2026, 11, 15)), ['1m', '2w', 'd2026-11-15'])
check("期限を過ぎたら知らせない", C.due_stages(_E, date(2026, 11, 16)), [])
check("1か月以内（期限切れを含む）", (C.in_window(_E, date(2026, 10, 14)), C.in_window(_E, date(2026, 10, 15)),
                                C.in_window(_E, date(2027, 1, 1))), (False, True, True))

# ===== 記録のキー =====
print("\n▼ 記録のキー")
_a = C.Account(2, '山田 太郎', date(2026, 11, 3), '2026年11月03日', '1427112485047242945', 'ok', '1427112485047242945')
_n = C._notice('1m', _a, discord_id=_a.discord_id)
check("キーは64文字のハッシュ", len(_n['key']), 64)
check("記録する項目", {k: v for k, v in _n.items() if k != 'key'},
      {'kind': '1m', 'account_name': '山田 太郎', 'expiry_date': date(2026, 11, 3), 'discord_id': '1427112485047242945'})
check("同じ内容なら同じキー", C._notice('1m', _a, discord_id=_a.discord_id)['key'], _n['key'])
check("区切りが違えば別のキー", C._notice('2w', _a, discord_id=_a.discord_id)['key'] != _n['key'], True)
check("宛先のIDが変われば別のキー（後から記入した人にも届く）",
      C._notice('1m', _a, discord_id='2222222222222222222')['key'] != _n['key'], True)
_renewed = C.Account(2, '山田 太郎', date(2027, 11, 3), '2027年11月03日', '1427112485047242945', 'ok', '')
check("有効期限が変われば別のキー（数え直し）", C._notice('1m', _renewed, discord_id=_a.discord_id)['key'] != _n['key'], True)
_spaced = C.Account(9, '山田　太郎', date(2026, 11, 3), '', '1427112485047242945', 'ok', '')
check("名義の空白の違いは同じキー", C._notice('1m', _spaced, discord_id=_a.discord_id)['key'], _n['key'])

# ===== 文面 =====
print("\n▼ 文面")
check("本人へのDM", C.build_personal_message(_a, date(2026, 10, 27)),
      "コート予約システムの利用登録の有効期限が近づいています。\n"
      "名義: 山田 太郎\n"
      "有効期限: 2026年11月03日（あと 7 日）\n"
      "期限までに更新をお願いします。")
check("期限日の当日は「本日まで」", "有効期限: 2026年11月03日（本日まで）" in C.build_personal_message(_a, date(2026, 11, 3)), True)

_list_accounts = [
    C.Account(2, '山田 太郎', date(2026, 11, 3), '2026年11月03日', '1427112485047242945', 'ok', ''),
    C.Account(3, '佐藤 花子', date(2026, 11, 15), '2026年11月15日', None, 'missing', ''),
    C.Account(4, '鈴木 一郎', date(2026, 10, 1), '2026年10月01日', '1111111111111111111', 'ok', ''),
    C.Account(5, '伊藤 四郎', date(2026, 12, 1), '2026年12月01日', None, 'missing', ''),
]
check("管理者への一覧（期限の近い順。1か月より先の行は載せない）",
      C.build_admin_message(_list_accounts, date(2026, 10, 16), {2: '本人に DM 済み', 3: 'Discord 未記入'}),
      "利用登録の有効期限が 1 か月以内のアカウント\n"
      "・鈴木 一郎　2026年10月01日（期限切れ）\n"
      "・山田 太郎　2026年11月03日（あと 18 日・本人に DM 済み）\n"
      "・佐藤 花子　2026年11月15日（あと 30 日・Discord 未記入）")
_unreadable = [
    C.Account(6, '田中 次郎', None, '2026/11/15', None, 'missing', ''),
    C.Account(7, '高橋 三郎', date(2027, 3, 1), '2027-03-01 00:00:00', None, 'invalid', '1.42711248504724E+18'),
]
check("読めなかった行は一覧の下に並べる",
      C.build_admin_message(_unreadable, date(2026, 10, 16), {}),
      "利用登録の有効期限が 1 か月以内のアカウント\n"
      "（なし）\n"
      "\n"
      "表で読めなかった行\n"
      "・田中 次郎（6行目）　有効期限「2026/11/15」が読めません\n"
      "・高橋 三郎（7行目）　Discord ID が読めません")

# ===== 発行SQL =====
print("\n▼ 発行SQL")


class _FakeCursor:
    def __init__(self, rows=None, rowcount=0, error=None):
        self.rows = rows or []
        self.rowcount = rowcount
        self.error = error
        self.executed = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    async def execute(self, sql, params=None):
        self.executed.append((' '.join(sql.split()), params))
        if self.error:
            raise self.error

    async def fetchall(self):
        return self.rows


class _FakeConn:
    def __init__(self, cursor):
        self._cursor = cursor

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    def cursor(self, *a, **kw):
        return self._cursor


class _FakePool:
    def __init__(self, cursor):
        self._cursor = cursor

    def acquire(self):
        return _FakeConn(self._cursor)


def _with_cursor(cursor, fn):
    orig = C.db
    C.db = types.SimpleNamespace(pool=_FakePool(cursor))
    try:
        return _run(fn())
    finally:
        C.db = orig


_stamp = datetime(2026, 10, 16, 9, 0, 0)

C._table_ready = False
_cur = _FakeCursor()
_with_cursor(_cur, lambda: C._ensure_table())
check("テーブルが無ければ作る", _cur.executed[0][0].startswith(
    "CREATE TABLE IF NOT EXISTS court_registration_expiry_notices"), True)
check("キーは一意（INSERT IGNORE で二重送信を防ぐため）", "UNIQUE KEY uq_notice_key (notice_key)" in _cur.executed[0][0], True)
_cur = _FakeCursor()
_with_cursor(_cur, lambda: C._ensure_table())
check("2回目以降は作らない", _cur.executed, [])

_cur = _FakeCursor(rows=[{'notice_key': 'k1'}])
check("記録済みのキー", _with_cursor(_cur, lambda: C._recorded_keys(['k1', 'k2'])), {'k1'})
check("記録済みのキーのSQL", _cur.executed[0],
      ("SELECT notice_key FROM court_registration_expiry_notices WHERE notice_key IN (%s, %s)", ['k1', 'k2']))
check("キー指定なしはDBを見ない", _with_cursor(_FakeCursor(error=Exception('x')), lambda: C._recorded_keys([])), set())

_cur = _FakeCursor(rowcount=1)
check("記録できれば True", _with_cursor(_cur, lambda: C._record_notice(_n, _stamp)), True)
check("記録は INSERT IGNORE", _cur.executed[0][0],
      "INSERT IGNORE INTO court_registration_expiry_notices "
      "(notice_key, kind, account_name, expiry_date, discord_id, notified_at) VALUES (%s, %s, %s, %s, %s, %s)")
check("記録のパラメータ", _cur.executed[0][1],
      (_n['key'], '1m', '山田 太郎', date(2026, 11, 3), '1427112485047242945', _stamp))
check("既に記録があれば False", _with_cursor(_FakeCursor(rowcount=0), lambda: C._record_notice(_n, _stamp)), False)
_cur = _FakeCursor(rowcount=1)
_with_cursor(_cur, lambda: C._record_notice({**_n, 'account_name': 'あ' * 300}, _stamp))
check("名義は255文字に収める", len(_cur.executed[0][1][2]), 255)

_cur = _FakeCursor(rowcount=1)
check("巻き戻し成功", _with_cursor(_cur, lambda: C._release_notice('k1', _stamp)), True)
check("巻き戻しは自分が書いた行だけ", _cur.executed[0],
      ("DELETE FROM court_registration_expiry_notices WHERE notice_key = %s AND notified_at = %s", ('k1', _stamp)))
check("巻き戻しのDBエラーは握って False",
      _with_cursor(_FakeCursor(error=Exception(2013, 'Lost connection')), lambda: C._release_notice('k1', _stamp)), False)
check("記録日時は秒精度", C._now_jst().microsecond, 0)
check("今日の日付はJST", C._today_jst(), datetime.now(C.JST).date())

# ===== Discord への送信 =====
print("\n▼ Discord への送信")


class _FakeResponse:
    def __init__(self, status, payload=None, text='', headers=None):
        self.status_code = status
        self._payload = payload
        self.text = text
        self.headers = headers or {}

    def json(self):
        return self._payload


class _FakeClient:
    def __init__(self, responses=None, error=None):
        self.responses = list(responses or [])
        self.error = error
        self.posts = []

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return False

    async def post(self, url, headers=None, json=None, timeout=None):
        self.posts.append((url, headers, json))
        if self.error:
            raise self.error
        return self.responses.pop(0)


def _with_token(token, fn):
    saved = os.environ.get('DISCORD_BOT_TOKEN')
    if token is None:
        os.environ.pop('DISCORD_BOT_TOKEN', None)
    else:
        os.environ['DISCORD_BOT_TOKEN'] = token
    try:
        return _run(fn())
    finally:
        if saved is None:
            os.environ.pop('DISCORD_BOT_TOKEN', None)
        else:
            os.environ['DISCORD_BOT_TOKEN'] = saved


_cl = _FakeClient([_FakeResponse(200, {'id': '555'}), _FakeResponse(200)])
check("成功", _with_token('tkn', lambda: C._send_dm(_cl, '1427112485047242945', '本文')), (True, ''))
check("DMチャンネルを作る", (_cl.posts[0][0], _cl.posts[0][2]),
      ('https://discord.com/api/v10/users/@me/channels', {'recipient_id': '1427112485047242945'}))
check("Bot トークンで送る", _cl.posts[0][1]['Authorization'], 'Bot tkn')
check("作ったチャンネルへ送る（メンションは発火させない）", (_cl.posts[1][0], _cl.posts[1][2]),
      ('https://discord.com/api/v10/channels/555/messages', {'content': '本文', 'allowed_mentions': {'parse': []}}))
check("トークン未設定は送らない", _with_token(None, lambda: C._send_dm(_FakeClient(), '1', 'x')),
      (False, 'DISCORD_BOT_TOKEN未設定'))
_ok, _reason = _with_token('tkn', lambda: C._send_dm(_FakeClient([_FakeResponse(403, text='Missing Access')]), '1', 'x'))
check("DMチャンネルを作れなければ失敗", (_ok, 'DMチャンネル作成 status=403' in _reason), (False, True))
_ok, _reason = _with_token('tkn', lambda: C._send_dm(
    _FakeClient([_FakeResponse(200, {'id': '555'}), _FakeResponse(403, text='Cannot send messages to this user')]),
    '1', 'x'))
check("DMを受け付けない相手は失敗", (_ok, 'status=403' in _reason), (False, True))
check("例外は失敗として返す",
      _with_token('tkn', lambda: C._send_dm(_FakeClient(error=Exception('boom')), '1', 'x'))[0], False)
_cl = _FakeClient([_FakeResponse(200, {'id': '555'})] + [_FakeResponse(200)] * 6)
_long = '\n'.join(f'・名義{i}　2026年11月{i % 28 + 1:02d}日（あと {i} 日・Discord 未記入）' for i in range(120))
check("長い本文も送れる", _with_token('tkn', lambda: C._send_dm(_cl, '1', _long))[0], True)
check("長い本文は2000文字以内に分けて送る",
      (len(_cl.posts) > 2, all(len(p[2]['content']) <= 2000 for p in _cl.posts[1:])), (True, True))
_orig_asyncio = C.asyncio
C.asyncio = types.SimpleNamespace(sleep=lambda s: _orig_asyncio.sleep(0))
try:
    _cl = _FakeClient([_FakeResponse(429, headers={'Retry-After': '0.1'}), _FakeResponse(200, {'id': '555'}),
                       _FakeResponse(200)])
    check("429 は1度だけ待って再送", (_with_token('tkn', lambda: C._send_dm(_cl, '1', 'x')), len(_cl.posts)),
          ((True, ''), 3))
finally:
    C.asyncio = _orig_asyncio

# ===== エンドポイントの流れ =====
print("\n▼ エンドポイント（notify_court_registration_expiry）")


class _FakeStore:
    """DB・Discord・表の取得を差し替えて、エンドポイントを日ごとに呼ぶ"""

    def __init__(self, fail_ids=(), release_fail=False, record_error_ids=(), table_error=None,
                 fetch_error=None, taken_by_other=False):
        self.recorded = {}      # key -> kind
        self.fail_ids = set(fail_ids)
        self.release_fail = release_fail
        self.record_error_ids = set(record_error_ids)
        self.table_error = table_error
        self.fetch_error = fetch_error
        self.taken_by_other = taken_by_other
        self.sent = []          # (discord_id, content)
        self.releases = []
        self.sleeps = 0

    async def ensure_table(self):
        if self.table_error:
            raise self.table_error

    async def recorded_keys(self, keys):
        return {k for k in keys if k in self.recorded}

    async def record(self, notice, stamp):
        if notice['discord_id'] in self.record_error_ids:
            raise Exception(2013, 'Lost connection')
        if self.taken_by_other or notice['key'] in self.recorded:
            return False
        self.recorded[notice['key']] = notice['kind']
        return True

    async def release(self, key, stamp):
        self.releases.append(key)
        if self.release_fail:
            return False
        self.recorded.pop(key, None)
        return True

    async def send(self, client, discord_id, content):
        self.sent.append((discord_id, content))
        return (False, 'status=403') if discord_id in self.fail_ids else (True, '')

    def personal(self):
        return [(d, c) for d, c in self.sent if d != ADMIN]

    def admin(self):
        return [c for d, c in self.sent if d == ADMIN]


ADMIN = C.DEFAULT_ADMIN_DISCORD_ID
_ROWS = [
    ('山田 太郎', 'ヤマダ タロウ', '2026年11月03日', '4444444444444444444'),   # 2行目
    ('佐藤 花子', 'サトウ ハナコ', '2026年11月15日', None),                    # 3行目 ID未記入
    ('鈴木 一郎', 'スズキ イチロウ', '2026年10月01日', '1111111111111111111'), # 4行目 期限切れ
    ('伊藤 四郎', 'イトウ シロウ', '2026年12月01日', '2222222222222222222'),   # 5行目 まだ先
    ('田中 次郎', 'タナカ ジロウ', '2026/11/15', '3333333333333333333'),       # 6行目 有効期限が読めない
    ('高橋 三郎', 'タカハシ サブロウ', datetime(2027, 3, 1), 1.4271124850472428e18),  # 7行目 IDが数値
]


def _call(store, today, rows=None, dry_run=False, admin_env=None):
    """表 rows と今日の日付 today でエンドポイントを呼ぶ。admin_env は管理者の宛先の環境変数（None=未設定）"""
    names = ('_download_workbook', '_ensure_table', '_recorded_keys', '_record_notice', '_release_notice',
             '_send_dm', '_today_jst', 'httpx', 'asyncio')
    saved = {k: getattr(C, k) for k in names}
    saved_env = os.environ.get('COURT_EXPIRY_ADMIN_DISCORD_ID')
    if admin_env is None:
        os.environ.pop('COURT_EXPIRY_ADMIN_DISCORD_ID', None)
    else:
        os.environ['COURT_EXPIRY_ADMIN_DISCORD_ID'] = admin_env

    async def _download():
        if store.fetch_error:
            raise store.fetch_error
        return _workbook(_ROWS if rows is None else rows)

    async def _sleep(sec):
        store.sleeps += 1

    C._download_workbook = _download
    C._ensure_table = store.ensure_table
    C._recorded_keys = store.recorded_keys
    C._record_notice = store.record
    C._release_notice = store.release
    C._send_dm = store.send
    C._today_jst = lambda: today
    C.httpx = types.SimpleNamespace(AsyncClient=lambda: _FakeClient())
    C.asyncio = types.SimpleNamespace(sleep=_sleep)
    try:
        return _run(C.notify_court_registration_expiry(dry_run=dry_run))
    except C.HTTPException as e:
        return {'http_error': e.status_code, 'detail': e.detail}
    finally:
        for k, v in saved.items():
            setattr(C, k, v)
        if saved_env is None:
            os.environ.pop('COURT_EXPIRY_ADMIN_DISCORD_ID', None)
        else:
            os.environ['COURT_EXPIRY_ADMIN_DISCORD_ID'] = saved_env


def _personal_ids(store, start=0):
    return [d for d, _ in store.personal()[start:]]


_D = date(2026, 10, 16)
_st = _FakeStore()
_res = _call(_st, _D)
check("初回: 1か月以内に入っている人に送る（期限切れ・まだ先の人には送らない）", _personal_ids(_st), ['4444444444444444444'])
check("初回: 本人へのDMの文面", _st.personal()[0][1],
      "コート予約システムの利用登録の有効期限が近づいています。\n"
      "名義: 山田 太郎\n"
      "有効期限: 2026年11月03日（あと 18 日）\n"
      "期限までに更新をお願いします。")
check("初回: 結果", (_res['sent_count'], _res['failed_count'], _res['results'][0]['stages']), (1, 0, ['1m']))
check("初回: 管理者へ一覧を送る", _st.admin(), [
    "利用登録の有効期限が 1 か月以内のアカウント\n"
    "・鈴木 一郎　2026年10月01日（期限切れ）\n"
    "・山田 太郎　2026年11月03日（あと 18 日・本人に DM 済み）\n"
    "・佐藤 花子　2026年11月15日（あと 30 日・Discord 未記入）\n"
    "\n"
    "表で読めなかった行\n"
    "・田中 次郎（6行目）　有効期限「2026/11/15」が読めません\n"
    "・高橋 三郎（7行目）　Discord ID が読めません"])
check("初回: 一覧のきっかけ（新しく1か月以内に入った2人と、読めない行2つ）",
      sorted(_res['admin']['triggers']), ['bad_expiry', 'bad_id', 'entered', 'entered'])
check("初回: 管理者への一覧の前は間隔を空ける", _st.sleeps, 1)

_res = _call(_st, _D)
check("同じ日にもう一度動いても送らない", (len(_st.sent), _res['sent_count'], _res['admin']), (2, 0, None))
_res = _call(_st, date(2026, 10, 17))
check("翌日は何も送らない", (len(_st.sent), _res['admin']), (2, None))
_res = _call(_st, date(2026, 10, 20))
check("2週間前に本人へ送る（一覧は送らない）", (_personal_ids(_st, 1), _res['results'][0]['stages'], _res['admin']),
      (['4444444444444444444'], ['2w'], None))
_call(_st, date(2026, 10, 26))
check("8日前は送らない", len(_st.personal()), 2)
_call(_st, date(2026, 10, 27))
_call(_st, date(2026, 10, 28))
check("1週間前から毎日送る", (len(_st.personal()), 'あと 6 日' in _st.personal()[-1][1]), (4, True))
_res = _call(_st, date(2026, 11, 1))
check("11/1: 山田さんの毎日分と、新しく1か月以内に入った伊藤さん",
      _personal_ids(_st, 4), ['4444444444444444444', '2222222222222222222'])
check("11/1: 伊藤さんが新しく入ったので一覧を送る", _st.admin()[-1].split('\n')[:5], [
    "利用登録の有効期限が 1 か月以内のアカウント",
    "・鈴木 一郎　2026年10月01日（期限切れ）",
    "・山田 太郎　2026年11月03日（あと 2 日・本人に DM 済み）",
    "・佐藤 花子　2026年11月15日（あと 14 日・Discord 未記入）",
    "・伊藤 四郎　2026年12月01日（あと 30 日・本人に DM 済み）"])
check("11/1: 一覧のきっかけは伊藤さんだけ", _res['admin']['triggers'], ['entered'])
_call(_st, date(2026, 11, 3))
check("期限日の当日は「本日まで」", "（本日まで）" in _st.personal()[-1][1], True)
_n_before = len(_st.sent)
_call(_st, date(2026, 11, 4))
check("期限を過ぎたら本人には送らない", len(_st.sent), _n_before)

_renewed_rows = [('山田 太郎', 'ヤマダ タロウ', '2027年11月03日', '4444444444444444444')] + _ROWS[1:]
_n_before = len(_st.sent)
_call(_st, date(2026, 11, 5), rows=_renewed_rows)
check("表の有効期限が書き換わったら止まる", len(_st.sent), _n_before)
_res = _call(_st, date(2027, 10, 3), rows=_renewed_rows)
check("新しい期限の1か月前から数え直す",
      ('4444444444444444444' in [d for d, _ in _st.sent[_n_before:]], 'entered' in _res['admin']['triggers']),
      (True, True))

print("\n▼ 取りこぼし・後からの記入")
_st = _FakeStore()
_res = _call(_st, date(2026, 10, 24))
check("初めて見たのが10日前なら1通にまとめる（1か月前と2週間前を記録）",
      (_personal_ids(_st), _res['results'][0]['stages']), (['4444444444444444444'], ['1m', '2w']))
_st = _FakeStore()
_call(_st, date(2026, 10, 16))
_filled = list(_ROWS)
_filled[1] = ('佐藤 花子', 'サトウ ハナコ', '2026年11月15日', '5555555555555555555')
_res = _call(_st, date(2026, 10, 20), rows=_filled)
check("IDを後から記入した人には、その時点の区切りで送る", '5555555555555555555' in _personal_ids(_st), True)
check("後から記入しても一覧は送り直さない", _res['admin'], None)

print("\n▼ 失敗のとき")
_st = _FakeStore(fail_ids=['4444444444444444444'])
_res = _call(_st, _D)
check("本人へ送れなければ失敗に数える", (_res['sent_count'], _res['failed_count'], _res['results'][0]['status']),
      (0, 1, 'failed'))
check("本人へ送れなければ記録を戻す（次回送り直す）", len(_st.releases), 1)
check("一覧には「DM が届かず」と出す", "・山田 太郎　2026年11月03日（あと 18 日・DM が届かず）" in _st.admin()[0], True)
_st.fail_ids = set()
_call(_st, date(2026, 10, 17))
check("翌日に送り直す", _personal_ids(_st), ['4444444444444444444', '4444444444444444444'])

_st = _FakeStore(fail_ids=[ADMIN])
_res = _call(_st, _D)
check("管理者へ送れなければ失敗に数える", (_res['admin']['status'], _res['failed_count']), ('failed', 1))
check("管理者へ送れなければきっかけの記録を戻す", len(_st.releases), 4)
_st.fail_ids = set()
_res = _call(_st, date(2026, 10, 17))
check("翌日に一覧を送り直す", (_res['admin']['status'], len(_st.admin())), ('sent', 2))

_st = _FakeStore(fail_ids=['4444444444444444444'], release_fail=True)
_res = _call(_st, _D)
check("記録を戻せないときは理由に出す", '記録を戻せませんでした' in _res['results'][0]['reason'], True)

_st = _FakeStore(record_error_ids=['4444444444444444444'])
_res = _call(_st, _D)
check("記録のDBエラーは失敗にして送らない",
      (_res['results'][0]['status'], _res['results'][0]['reason'].startswith('db:'), _personal_ids(_st)),
      ('failed', True, []))
check("記録のDBエラーでも一覧は送る", len(_st.admin()), 1)

_st = _FakeStore(taken_by_other=True)
_res = _call(_st, _D)
check("並行して走った実行が先に記録していたら送らない", (_st.sent, _res['results'], _res['admin']), ([], [], None))

_st = _FakeStore(fetch_error=Exception('timeout'))
_res = _call(_st, _D)
check("表の取得失敗は502（記録も送信もしない）", (_res.get('http_error'), _st.recorded, _st.sent), (502, {}, []))
_st = _FakeStore(table_error=Exception(1142, 'CREATE command denied'))
_res = _call(_st, _D)
check("記録を参照できなければ500（送らない）", (_res.get('http_error'), _st.sent), (500, []))

print("\n▼ 管理者の宛先")
_st = _FakeStore()
_call(_st, _D, admin_env='9999999999999999999')
check("環境変数で宛先を差し替えられる", [d for d, _ in _st.sent][-1], '9999999999999999999')
_st = _FakeStore()
_res = _call(_st, _D, admin_env='')
check("宛先が空なら一覧は送らない（本人には送る）",
      (_res['admin']['status'], _personal_ids(_st), _st.admin()), ('skipped', ['4444444444444444444'], []))
check("宛先が空ならきっかけを記録しない（設定後に送れるように）",
      sorted(_st.recorded.values()), ['1m'])

print("\n▼ dry_run")
_st = _FakeStore()
_res = _call(_st, _D, dry_run=True)
check("dry_run は記録も送信もしない", (_st.recorded, _st.sent), ({}, []))
check("dry_run は本人へ送る予定を返す",
      [(p['name'], p['discord_id'], p['stages']) for p in _res['personal']],
      [('山田 太郎', '4444444444444444444', ['1m'])])
check("dry_run は送る予定の文面を返す", _res['personal'][0]['message'], C.build_personal_message(
    C.Account(2, '山田 太郎', date(2026, 11, 3), '', '4444444444444444444', 'ok', ''), _D))
check("dry_run は一覧の予定を返す",
      (_res['admin']['to'], "・山田 太郎　2026年11月03日（あと 18 日・本人に DM 予定）" in _res['admin']['message']),
      (ADMIN, True))
_call(_st, _D)
_res = _call(_st, _D, dry_run=True)
check("送ったあとの dry_run は予定なし", (_res['personal'], _res['admin']), ([], None))

# ===== 設定 =====
print("\n▼ 設定")
check("表の列（G=氏名, H=カタカナ, I=有効期限, M=Discord ID）",
      (C.NAME_COL, C.KANA_COL, C.EXPIRY_COL, C.DISCORD_ID_COL), (6, 7, 8, 12))
check("区切り（2週間前・1週間前から毎日）", (C.TWO_WEEKS_DAYS, C.DAILY_FROM_DAYS), (14, 7))
check("JSTは+9時間固定", C.JST.utcoffset(None) == timedelta(hours=9), True)
check("エンドポイント", [(r.path, sorted(r.methods)) for r in C.router.routes],
      [('/notify/court-registration-expiry', ['POST'])])
_script = (Path(__file__).parent / 'scripts' / 'notify_court_registration_expiry.sh').read_bytes()
check("cron 用スクリプトが同じエンドポイントを呼ぶ",
      b'http://localhost:8000/api/notify/court-registration-expiry' in _script, True)
check("cron 用スクリプトは失敗件数を見る", b'"failed_count": *[1-9]' in _script, True)

print()
if _failures:
    print(f"❌ {len(_failures)}件失敗")
    for f in _failures:
        print(f"  - {f}")
    sys.exit(1)
print("✅ すべて成功")
